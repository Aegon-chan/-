# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path

out = (Path(__file__).parent /
       "青少年山地科学营2026_报告厅流程PPT模板_20页可编辑版.pptx")
assets = Path(__file__).parent / "assets"
bg = assets / "cover_bg.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

pages = [
    ("探秘山地力量\n青少年山地科学营2026", "活动主视觉封面"),
    ("开营仪式", "探索山地科学之旅正式启程"),
    ("开营仪式流程", "领导致辞\u3001专家寄语\u3001授旗仪式\u3001学员发言\u3001宣布开营"),
    ("领导致辞", "LEADERS' MESSAGE"),
    ("领导嘉宾介绍", "姓名\u3001职务\u3001单位\u3001寄语"),
    ("专家讲座", "聆听科学声音 探索山地奥秘"),
    ("专家介绍", "专家照片\u3001研究方向\u3001代表成果"),
    ("讲座主题", "主讲专家:XXX教授\n主题:山地科学探索"),
    ("科学实践探索", "从课堂走向山地现场"),
    ("实践任务发布", "地貌观察\u3001野外调查\u3001科学记录\u3001成果汇报"),
    ("结营仪式", "山海相遇 科学同行"),
    ("荣誉时刻", "表彰优秀探索成果"),
    ("优秀营员奖", "获奖名单展示"),
    ("优秀团队展示", "团队名称\u3001成员信息"),
    ("活动成果回顾", "课堂学习\u3001野外实践\u3001专家交流"),
    ("营员感言", "少年的科学梦想"),
    ("成果展示", "探索成果与精彩瞬间"),
    ("感谢同行", "感谢专家\u3001老师和每一位营员"),
    ("毕营总结", "探索不止步 山地未来见"),
    ("结束页", "青少年山地科学营2026\n中国\u00b7成都")
]

for idx, (title, subtitle) in enumerate(pages):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(bg), 0, 0,
                             width=prs.slide_width,
                             height=prs.slide_height)

    t = slide.shapes.add_textbox(Inches(0.8), Inches(1.0),
                                 Inches(9.5), Inches(1.4))
    p = t.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(34 if idx == 0 else 30)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    s = slide.shapes.add_textbox(Inches(0.85), Inches(2.7),
                                 Inches(9.5), Inches(1))
    p2 = s.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(240, 245, 240)

    f = slide.shapes.add_textbox(Inches(0.8), Inches(6.8),
                                 Inches(5), Inches(0.3))
    pf = f.text_frame.paragraphs[0]
    pf.text = "青少年山地科学营2026  |  {:02d}".format(idx + 1)
    pf.font.size = Pt(10)
    pf.font.color.rgb = RGBColor(230, 240, 235)

prs.save(out)
print("OK ->", out)
