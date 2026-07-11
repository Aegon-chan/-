"""
文件上传 API

提供文件上传端点，将文件保存到 data/uploads/ 目录，
返回可通过静态挂载路径访问的 URL。
"""

from __future__ import annotations

import uuid
import os
import mimetypes
from pathlib import Path

from fastapi import APIRouter, UploadFile, File, HTTPException

router = APIRouter()

UPLOAD_DIR = Path(__file__).resolve().parents[1] / "data" / "uploads"


@router.post("/upload", summary="Upload file")
async def upload_file(file: UploadFile = File(...)):
    """上传文件并返回可访问的 URL"""
    if not file.filename:
        raise HTTPException(status_code=400, detail="Filename cannot be empty")

    # 确保上传目录存在
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

    # 生成唯一文件名，保留扩展名
    ext = os.path.splitext(file.filename)[1] or ""
    unique_name = f"{uuid.uuid4().hex}{ext}"
    save_path = UPLOAD_DIR / unique_name

    # 保存文件
    content = await file.read()
    save_path.write_bytes(content)

    # 获取 MIME 类型
    mime_type, _ = mimetypes.guess_type(file.filename)
    if mime_type is None:
        mime_type = "application/octet-stream"

    return {
        "url": f"/uploads/{unique_name}",
        "filename": file.filename,
        "size": len(content),
        "mime_type": mime_type,
    }

# ───────────────────────────── 文件处理工具函数 ─────────────────────────────


def extract_text_from_file(file_path: str, mime_type: str) -> str:
    """从文件中提取文本内容（支持 PDF / DOCX / TXT / MD）"""
    import logging
    logger = logging.getLogger("research_agent.uploads")

    fp = Path(file_path)
    if not fp.exists():
        logger.warning("File not found: %s", file_path)
        return ""

    try:
        # TXT / MD / 纯文本
        if mime_type in ("text/plain", "text/markdown", "text/csv", "application/json"):
            return fp.read_text(encoding="utf-8", errors="replace")

        # PDF
        if mime_type == "application/pdf":
            try:
                import fitz  # PyMuPDF
                text_parts = []
                doc = fitz.open(file_path)
                for page in doc:
                    text_parts.append(page.get_text())
                doc.close()
                combined = "\n".join(text_parts)
                return combined if combined.strip() else "[PDF 文件未能提取到文本内容]"
            except Exception as pdf_err:
                logger.warning("PyMuPDF extraction failed: %s", pdf_err)
                return "[PDF 文件提取失败]"

        # DOCX
        if mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",):
            try:
                from docx import Document
                doc = Document(file_path)
                paras = [p.text for p in doc.paragraphs]
                combined = "\n".join(paras)
                return combined if combined.strip() else "[DOCX 文件未能提取到文本内容]"
            except Exception as docx_err:
                logger.warning("DOCX extraction failed: %s", docx_err)
                return "[DOCX 文件提取失败]"

        # Excel (.xlsx / .xls)
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        ) or file_path.lower().endswith((".xlsx", ".xls", ".xlsm", ".xlsb")):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(file_path, read_only=True, data_only=True)
                text_parts = []
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        row_values = [str(cell).strip() if cell is not None else "" for cell in row]
                        filtered = [v for v in row_values if v]
                        if filtered:
                            rows.append(" | ".join(filtered))
                    if rows:
                        text_parts.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
                    else:
                        text_parts.append(f"--- Sheet: {sheet_name} ---\n(empty)")
                wb.close()
                combined = "\n\n".join(text_parts)
                return combined if combined.strip() else "[Excel file appears to be empty]"
            except Exception as xlsx_err:
                logger.warning("Excel extraction failed: %s", xlsx_err)
                try:
                    import csv
                    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                        reader = csv.reader(f)
                        rows = [" | ".join(row) for row in reader if any(cell.strip() for cell in row)]
                    return "\n".join(rows) if rows else "[Excel file could not be parsed]"
                except Exception:
                    return "[Excel file extraction failed]"

        logger.info("Unsupported mime_type for text extraction: %s", mime_type)
        return ""
    except Exception as e:
        logger.exception("extract_text_from_file error: %s", e)
        return ""


def image_to_data_uri(file_path: str) -> str:
    """将图片文件转为 data URI（base64 编码）"""
    import base64
    import mimetypes

    fp = Path(file_path)
    if not fp.exists():
        return ""

    raw = fp.read_bytes()
    encoded = base64.b64encode(raw).decode("ascii")
    mime, _ = mimetypes.guess_type(fp.name)
    if mime is None:
        mime = "image/png"
    return f"data:{mime};base64,{encoded}"

def process_attachments(attachments: list | None) -> "list[ContentPart]":
    """将附件元数据列表转为 ContentPart 列表，供多模态消息使用。

    支持 dict 格式（来自前端 _pending_attachments）和 AttachmentSchema 对象（来自 API）。
    图片 -> base64 data URI -> ContentPart(type="image_base64")
    文档（PDF/DOCX/TXT/MD）-> 文本提取 -> ContentPart(type="text")
    """
    from backend.models.base import ContentPart
    from pathlib import Path

    if not attachments:
        return []
    parts: list[ContentPart] = []
    upload_root = Path("data") / "uploads"
    for att in attachments:
        # Support both dict and Pydantic model
        if isinstance(att, dict):
            url = att.get("url", "")
            mime = att.get("mime_type", "")
        else:
            url = att.url
            mime = att.mime_type
        fname = Path(url).name
        file_path = upload_root / fname
        if not file_path.exists():
            continue
        try:
            if mime.startswith("image/"):
                data_uri = image_to_data_uri(str(file_path))
                if data_uri:
                    parts.append(ContentPart(type="image_base64", image_base64=data_uri, mime_type=mime))
            else:
                text = extract_text_from_file(str(file_path), mime)
                if text:
                    parts.append(ContentPart(type="text", text=text))
        except Exception:
            continue
    return parts
