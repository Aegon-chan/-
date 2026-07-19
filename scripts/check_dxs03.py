import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\24622\Desktop\监测数据2020.5.30-2026.6.10.xlsx", data_only=True)
ws = wb["监测数据合集"]

# Check col headers
print("Excel column headers (row 1):")
for c in [1, 3, 4, 5]:
    print(f"  Col {c}: {ws.cell(1, c).value}")

print("\nRow 5 notes:")
for c in [3, 4, 5, 6]:
    v = ws.cell(5, c).value
    if v: print(f"  Col {c}: {v}")

print("\nDXS3 (col 5) first 5 non-empty values:")
cnt = 0
for r in range(2, 2204):
    v = ws.cell(r, 5).value
    if v is not None and str(v).strip() != '':
        t = ws.cell(r, 1).value
        print(f"  Row {r}: {t} -> {v}")
        cnt += 1
        if cnt >= 5: break

print(f"\nDXS3 total non-empty cells: {sum(1 for r in range(2, 2204) if ws.cell(r, 5).value is not None and str(ws.cell(r,5).value).strip()!='')}")

print("\nComparison DXS1 vs DXS3 at key times:")
for r in [1500, 1800, 2000, 2100, 2200]:
    t = ws.cell(r, 1).value
    d1 = ws.cell(r, 3).value
    d3 = ws.cell(r, 5).value
    print(f"  Row {r} ({t}): DXS1={d1}, DXS3={d3}")

wb.close()
