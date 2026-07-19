from __future__ import annotations
import os, sys
from datetime import datetime
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import openpyxl
from sqlalchemy import create_engine, text, inspect as sa_inspect
from backend.db.monitoring_session import MonitoringBase
from backend.db import monitoring_models  # noqa: F401

EXCEL = r"C:\Users\24622\Desktop\监测数据2020.5.30-2026.6.10.xlsx"
DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "monitoring_data.db")
engine = create_engine("sqlite:///" + DB.replace("\\", "/"), echo=False)
B = 500

def to_float(v):
    if v is None: return None
    try: return float(v)
    except: return None

def parse_time(v):
    if v is None: return None
    if isinstance(v, datetime): return v
    if isinstance(v, str):
        v = v.strip()
        if not v: return None
        for fmt in ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"]:
            try: return datetime.strptime(v, fmt)
            except: continue
    return None

stations = [
    ("RFS-01", "降雨量监测站", "rainfall", "日降雨量"),
    ("DXS01", "地下水位监测点1", "groundwater", "DXS1"),
    ("DXS02", "地下水位监测点2", "groundwater", "DXS2"),
    ("DXS03", "地下水位监测点3", "groundwater", "DXS3"),
    ("GNS01", "GNSS监测点G1", "gnss", "G1"),
    ("GNS02", "GNSS监测点G2", "gnss", "G2"),
    ("GNS03", "GNSS监测点G3", "gnss", "G3 (文章更名G4)"),
    ("GNS04", "GNSS监测点G4", "gnss", "G4"),
]
stype_of = {s[0]: s[2] for s in stations}

sections = [
    ("RFS-01", 1, [("rainfall_amount", 2)], ["rainfall_amount"]),
    ("DXS01", 1, [("water_level", 3)], ["water_level"]),
    ("DXS02", 1, [("water_level", 4)], ["water_level"]),
    ("DXS03", 1, [("water_level", 5)], ["water_level"]),
    ("GNS01", 7, [("dx",8),("dy",9),("dz",10),("planar_displacement",11),("variation",12),("daily_planar_rate",13),("daily_variation_rate",14)], ["dx","dy","dz"]),
    ("GNS02",16, [("dx",17),("dy",18),("dz",19),("planar_displacement",20),("variation",21),("daily_planar_rate",22),("daily_variation_rate",23)], ["dx","dy","dz"]),
    ("GNS03",25, [("dx",26),("dy",27),("dz",28),("planar_displacement",29),("variation",30),("daily_planar_rate",31),("daily_variation_rate",32)], ["dx","dy","dz"]),
    ("GNS04",34, [("dx",35),("dy",36),("dz",37),("planar_displacement",38),("variation",39),("daily_planar_rate",40),("daily_variation_rate",41)], ["dx","dy","dz"]),
]

tbl = {"rainfall":"rainfall_measurements","groundwater":"groundwater_measurements","gnss":"gnss_measurements"}

def run():
    for t in sa_inspect(engine).get_table_names():
        with engine.begin() as c: c.execute(text(f'DROP TABLE IF EXISTS "{t}"'))
    MonitoringBase.metadata.create_all(engine)
    print("Tables recreated")

    with engine.begin() as c:
        for code,name,stype,desc in stations:
            c.execute(text("INSERT INTO monitoring_stations (station_code,station_name,station_type,description,status) VALUES (:c,:n,:t,:d,'active')"),
                      {"c":code,"n":name,"t":stype,"d":desc})
    print("8 stations inserted")

    sid_map = {}
    with engine.connect() as c:
        for code,_,_,_ in stations:
            sid_map[code] = c.execute(text("SELECT id FROM monitoring_stations WHERE station_code=:c"), {"c": code}).scalar()

    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb["监测数据合集"]
    mr = ws.max_row; total = 0

    for scode, tcol, fields, req in sections:
        sid = sid_map[scode]
        table = tbl[stype_of[scode]]
        fnames = [x[0] for x in fields]
        req_set = set(req)
        all_c = ["station_id", "measurement_time"] + fnames
        rows = []

        for r in range(2, mr + 1):
            parsed_t = parse_time(ws.cell(r, tcol).value)
            if parsed_t is None: continue

            row = {"station_id": sid, "measurement_time": parsed_t}
            all_req_ok = True
            for fn, fc in fields:
                v = to_float(ws.cell(r, fc).value)
                row[fn] = v
                if fn in req_set and v is None:
                    all_req_ok = False

            if not all_req_ok: continue

            rows.append(row)
            if len(rows) >= B:
                bulk(engine, table, all_c, rows); total += len(rows); print(f"  {scode}: {total}..."); rows = []

        if rows: bulk(engine, table, all_c, rows); total += len(rows)
        print(f"  {scode}: done")

    print(f"\nTotal: {total} rows")
    wb.close()

def bulk(engine, table, cols, rows):
    sql = text(f"INSERT INTO {table} ({', '.join(f'\"{c}\"' for c in cols)}) VALUES ({', '.join(f':{c}' for c in cols)})")
    with engine.begin() as c: c.execute(sql, rows)

if __name__ == "__main__": run()
