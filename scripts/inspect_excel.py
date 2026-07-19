import openpyxl

path = r"C:\Users\24622\Desktop\监测数据2020.5.30-2026.6.10.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["监测数据合集"]

# Check G4 data existence and range
print("=== G4 data check ===")
for r in [2, 3, 100, 500, 1000, 1500, 2000, 2200, 2204]:
    vals = []
    for c in [25, 26, 34, 35]:  # G3 time, G3 dx, G4 time, G4 dx
        v = ws.cell(r, c).value
        vals.append(str(v)[:25] if v is not None else "")
    print(f"  Row {r}: G3(t={vals[0]}, dx={vals[1]})  G4(t={vals[2]}, dx={vals[3]})")

# Check DXS3 (col 5) vs DXS1 (col 3) - are they the same?
print("\n=== DXS1 vs DXS3 check ===")
for r in [100, 500, 1000, 1500, 2000, 2200, 2203]:
    d1 = ws.cell(r, 3).value
    d3 = ws.cell(r, 5).value
    print(f"  Row {r}: DXS1={d1}, DXS3={d3}")

# Check cols 33, 42-48 for any data
print("\n=== Trailing columns check ===")
for c in [33, 42, 43, 44, 45, 46, 47, 48]:
    vals = []
    for r in [1, 2, 3, 100, 500, 1000]:
        v = ws.cell(r, c).value
        if v is not None and v != "":
            vals.append(f"row{r}={str(v)[:20]}")
    if vals:
        print(f"  Col {c} ({str(ws.cell(1,c).value)[:30]}): {', '.join(vals)}")
    else:
        print(f"  Col {c} ({str(ws.cell(1,c).value)[:30]}): (empty)")

# Check last data row for G2 and G3 more precisely
print("\n=== Last data check for G2, G3, G4 ===")
for label, col_list in [("G2", [16, 17]), ("G3", [25, 26]), ("G4", [34, 35])]:
    last = 0
    for r in range(2, ws.max_row + 1):
        for c in col_list:
            if ws.cell(r, c).value is not None and ws.cell(r, c).value != "":
                last = r
                break
    vals = [(ws.cell(last, col_list[0]).value, ws.cell(last, col_list[1]).value)]
    print(f"  {label}: last row = {last}, vals = {[str(v)[:20] for v in vals[0]]}")

# Check if there are any G1-only rows (rows where G1 has data but others don't)
gns_only_count = 0
for r in range(2, ws.max_row + 1):
    g1 = ws.cell(r, 8).value
    g2 = ws.cell(r, 17).value
    g3 = ws.cell(r, 26).value
    g4 = ws.cell(r, 35).value
    if g1 is not None and g1 != "" and (g2 is None or g2 == "") and (g3 is None or g3 == "") and (g4 is None or g4 == ""):
        gns_only_count += 1

print(f"\nG1-only rows (G2/G3/G4 empty): {gns_only_count}")
